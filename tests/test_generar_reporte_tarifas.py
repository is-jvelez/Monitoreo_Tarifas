"""Tests de generar_reporte_tarifas.py. Correr desde la raiz del repo con:
    python -m unittest discover -s tests -v
"""

import contextlib
import io
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import generar_reporte_tarifas as g  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ENTRADA_OK = """\
Aduana
  [ELEGIBLE] DTYT_105853_2026 [tracking: No_intentos=3 -> se reseteara a 0]
  [ELEGIBLE] DTYT_102382_2025 (tarifa Estado=1-Inactivo, de todas formas se actualiza)
Resumen -> Evaluadas: 10 | Incompletas: 1 | Elegibles para UPDATE: 2

Contenedor
    [ELEGIBLE] DCCO_28421_2026
Resumen -> Evaluadas: 5 | Cliente anulado (excluidas): 4 | Elegibles para UPDATE: 1 (de las cuales con tarifa Estado=1-Inactivo: 0)

Carga Suelta
  (ninguna)
Resumen -> Evaluadas: 7 | Elegibles para UPDATE: 0

Servicios Adicionales
    [ELEGIBLE] DCSA_27000_2026 (ganadora de un grupo duplicado de 2 en el lote)
    [ELEGIBLE] DCSA_27001_2026
    [ELEGIBLE] DCSA_27002_2026
Resumen -> Evaluadas: 20 | Caso1 (ya en Dobra): 3 | Elegibles para UPDATE: 3
"""


def correr(*args):
    salida = io.StringIO()
    with contextlib.redirect_stdout(salida):
        codigo = g.main([str(a) for a in args])
    return codigo, salida.getvalue()


class TestParser(unittest.TestCase):
    def test_parsea_las_4_categorias_y_observaciones(self):
        resultado, avisos = g.parsear(ENTRADA_OK)
        self.assertEqual(avisos, [])
        self.assertEqual(g.validar(resultado), [])
        self.assertEqual(resultado["Aduana"]["codigos"], ["DTYT_105853_2026", "DTYT_102382_2025"])
        self.assertEqual(resultado["Aduana"]["observaciones"][0], "[tracking: No_intentos=3 -> se reseteara a 0]")
        self.assertEqual(resultado["Contenedor"]["observaciones"], [""])
        self.assertEqual(resultado["Carga Suelta"]["codigos"], [])  # asignada por eliminacion
        self.assertEqual(resultado["Carga Suelta"]["resumen"]["Evaluadas"], "7")
        self.assertEqual(len(resultado["Servicios Adicionales"]["codigos"]), 3)

    def test_conteo_que_no_calza(self):
        texto = ENTRADA_OK.replace("Elegibles para UPDATE: 3", "Elegibles para UPDATE: 4")
        problemas = g.validar(g.parsear(texto)[0])
        self.assertTrue(any("dice 4 elegibles, se parsearon 3" in p for p in problemas), problemas)

    def test_codigo_de_otra_categoria_en_el_bloque(self):
        texto = ENTRADA_OK.replace("DCSA_27001_2026", "DCCO_27001_2026")
        problemas = g.validar(g.parsear(texto)[0])
        self.assertTrue(any("Servicios Adicionales" in p and "DCCO_27001_2026" in p for p in problemas), problemas)

    def test_codigo_repetido(self):
        texto = ENTRADA_OK.replace("DCSA_27001_2026", "DCSA_27000_2026")
        problemas = g.validar(g.parsear(texto)[0])
        self.assertTrue(any("repetidos" in p for p in problemas), problemas)

    def test_codigos_sin_linea_resumen(self):
        texto = ENTRADA_OK + "    [ELEGIBLE] DCSA_99999_2026\n"
        _, avisos = g.parsear(texto)
        self.assertTrue(any("Se esperaban 4 bloques" in a for a in avisos), avisos)


class TestMain(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.entrada = self.tmp / "entrada.txt"
        self.entrada.write_text(ENTRADA_OK, encoding="utf-8")
        self.salida = self.tmp / "out"

    def tearDown(self):
        self._tmp.cleanup()

    def test_genera_excel_y_los_3_tipos_de_sql(self):
        codigo, _ = correr(self.entrada, self.salida, "--fecha", "2026-10-07")
        self.assertEqual(codigo, 0)
        for sub, nombre in (("Update", "script_Regularizacion"), ("update_lite", "script_Regularizacion"),
                            ("Queries", "script_consulta")):
            archivos = sorted(p.name for p in (self.salida / sub).glob("*.sql"))
            self.assertEqual(archivos, [f"{nombre}_{p}.sql" for p in ("DCCO", "DCSA", "DTYT")])

        update = (self.salida / "Update" / "script_Regularizacion_DTYT.sql").read_text(encoding="utf-8")
        self.assertIn("/* Ejecucion: 07/10/2026 | Tarifas a regularizar: 2 */", update)
        self.assertIn("    ('DTYT_105853_2026'),\n    ('DTYT_102382_2025');", update)

        lite = (self.salida / "update_lite" / "script_Regularizacion_DTYT.sql").read_text(encoding="utf-8")
        self.assertIn("'Regularizado manualmente - 20261007'", lite)
        self.assertEqual(lite.count("    'DTYT_105853_2026',\n    'DTYT_102382_2025'\n  );"), 2)

        consulta = (self.salida / "Queries" / "script_consulta_DCSA.sql").read_text(encoding="utf-8")
        self.assertIn("'DCSA_27000_2026',\n        'DCSA_27001_2026',\n        'DCSA_27002_2026'", consulta)

        for s in ("Update", "update_lite", "Queries"):
            for f in (self.salida / s).glob("*.sql"):
                texto = f.read_text(encoding="utf-8")
                self.assertNotIn("<lista_tarifas", texto, f)
                self.assertNotIn("{", texto, f)

        wb = load_workbook(self.salida / "Resumen_tarifas.xlsx")
        self.assertEqual(wb.sheetnames, ["Resumen general", "Aduana", "Carga Suelta", "Contenedor",
                                         "Servicios Adicionales"])
        self.assertEqual(wb["Resumen general"]["C5"].value, 10)  # Evaluadas de Aduana, como numero
        self.assertEqual(wb["Aduana"]["A5"].value, "DTYT_105853_2026")
        self.assertIn("No_intentos=3", wb["Aduana"]["B5"].value)

    def test_advertencia_bloquea_los_sql_pero_genera_excel(self):
        self.entrada.write_text(ENTRADA_OK.replace("Elegibles para UPDATE: 3", "Elegibles para UPDATE: 9"),
                                encoding="utf-8")
        codigo, texto = correr(self.entrada, self.salida)
        self.assertEqual(codigo, 1)
        self.assertIn("NO se generaron los .sql", texto)
        self.assertTrue((self.salida / "Resumen_tarifas.xlsx").exists())
        self.assertEqual(list(self.salida.rglob("*.sql")), [])

    def test_forzar_genera_los_sql_igual(self):
        self.entrada.write_text(ENTRADA_OK.replace("Elegibles para UPDATE: 3", "Elegibles para UPDATE: 9"),
                                encoding="utf-8")
        codigo, _ = correr(self.entrada, self.salida, "--forzar")
        self.assertEqual(codigo, 0)
        self.assertTrue((self.salida / "Update" / "script_Regularizacion_DCSA.sql").exists())

    def test_codigo_invalido_no_genera_sql_ni_con_forzar(self):
        self.entrada.write_text(ENTRADA_OK.replace("DCSA_27001_2026", "DCSA_27001_2026');DROP"),
                                encoding="utf-8")
        codigo, texto = correr(self.entrada, self.salida, "--forzar")
        self.assertEqual(codigo, 1)
        self.assertIn("codigos invalidos", texto)
        self.assertEqual(list(self.salida.rglob("*.sql")), [])

    def test_borra_sql_de_una_corrida_anterior(self):
        viejo = self.salida / "Update" / "script_Regularizacion_DCCS.sql"
        viejo.parent.mkdir(parents=True)
        viejo.write_text("UPDATE viejo", encoding="utf-8")
        otro = self.salida / "Update" / "notas.sql"  # no es generado por el script: no se toca
        otro.write_text("--", encoding="utf-8")

        codigo, texto = correr(self.entrada, self.salida)
        self.assertEqual(codigo, 0)
        self.assertFalse(viejo.exists())  # Carga Suelta hoy tiene 0 elegibles
        self.assertTrue(otro.exists())
        self.assertIn("se borraron 1 .sql", texto)

    def test_lee_entrada_utf16_de_ssms(self):
        self.entrada.write_text(ENTRADA_OK, encoding="utf-16")
        codigo, _ = correr(self.entrada, self.salida)
        self.assertEqual(codigo, 0)

    def test_carpeta_de_salida_por_defecto(self):
        cwd = os.getcwd()
        os.chdir(self.tmp)
        try:
            codigo, _ = correr(self.entrada, "--fecha", "2026-10-07")
        finally:
            os.chdir(cwd)
        self.assertEqual(codigo, 0)
        self.assertTrue((self.tmp / "Output_20261007" / "Resumen_tarifas.xlsx").exists())


class TestPlantillas(unittest.TestCase):
    def test_plantilla_sin_placeholder_falla_sin_escribir(self):
        resultado, _ = g.parsear(ENTRADA_OK)
        with tempfile.TemporaryDirectory() as tmp:
            for f in g.TEMPLATES_DIR.glob("*.sql"):
                (Path(tmp) / f.name).write_text(f.read_text(encoding="utf-8"), encoding="utf-8")
            roto = Path(tmp) / "update_DCCO.sql"
            roto.write_text(roto.read_text(encoding="utf-8").replace("    (<lista_tarifas>);", "(<lista_tarifas>);"),
                            encoding="utf-8")
            original = g.TEMPLATES_DIR
            g.TEMPLATES_DIR = Path(tmp)
            try:
                with self.assertRaisesRegex(ValueError, "update_DCCO.sql"):
                    g.construir_scripts_sql(resultado, date(2026, 10, 7))
            finally:
                g.TEMPLATES_DIR = original

    def test_todas_las_plantillas_reales_tienen_sus_placeholders(self):
        resultado = {cat: {"codigos": [f"{p}_1_2026"], "observaciones": [""], "resumen": {}}
                     for cat, p in g.CATEGORIAS.items()}
        scripts, _ = g.construir_scripts_sql(resultado, date(2026, 10, 7))
        self.assertEqual(len(scripts), 12)


if __name__ == "__main__":
    unittest.main()
